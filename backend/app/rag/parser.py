from pathlib import Path
import csv
import zipfile

import fitz
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation

class ParsedDocument(dict):
    pass


def _clean(text: str) -> str:
    return "\n".join(line.strip() for line in text.replace("\x00", "").splitlines() if line.strip())


def _parse_pptx(path: str) -> list[dict]:
    slides: list[dict] = []
    try:
        prs = Presentation(path)
        for index, slide in enumerate(prs.slides, 1):
            lines: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    lines.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
            notes = getattr(slide, "notes_slide", None)
            if notes and notes.notes_text_frame:
                lines.append(notes.notes_text_frame.text)
            text = _clean("\n".join(lines))
            if text:
                slides.append({"slide": index, "text": text})
    except Exception:
        # Last-resort fallback for unusual PPTX files: read embedded XML text nodes.
        with zipfile.ZipFile(path) as archive:
            slide_names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            for index, name in enumerate(slide_names, 1):
                xml = archive.read(name).decode("utf-8", errors="ignore")
                import re
                text = _clean("\n".join(re.findall(r"<a:t>(.*?)</a:t>", xml)))
                if text:
                    slides.append({"slide": index, "text": text})
    return slides


def parse_file(path: str) -> ParsedDocument:
    file_path = Path(path)
    ext = file_path.suffix.lower()
    pages: list[dict] = []
    if ext == ".pdf":
        with fitz.open(path) as doc:
            for index, page in enumerate(doc, 1):
                text = _clean(page.get_text("text"))
                if text:
                    pages.append({"page": index, "text": text})
    elif ext in [".docx", ".doc"]:
        doc = DocxDocument(path)
        text = _clean("\n".join(paragraph.text for paragraph in doc.paragraphs))
        pages.append({"section": "document", "text": text})
    elif ext in [".pptx", ".ppt"]:
        pages.extend(_parse_pptx(path))
    elif ext in [".xlsx", ".xls"]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.sheetnames:
            rows = []
            for row in workbook[sheet].iter_rows(values_only=True):
                rows.append(" | ".join("" if value is None else str(value) for value in row))
            text = _clean("\n".join(rows))
            if text:
                pages.append({"sheet": sheet, "text": text})
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="ignore") as handle:
            reader = csv.reader(handle)
            pages.append({"section": "csv", "text": _clean("\n".join(" | ".join(row) for row in reader))})
    elif ext in [".txt", ".md"]:
        pages.append({"section": ext[1:], "text": _clean(file_path.read_text(encoding="utf-8", errors="ignore"))})
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    return ParsedDocument(filename=file_path.name, pages=[page for page in pages if page.get("text", "").strip()])
