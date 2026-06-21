from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app


def test_health():
    with TestClient(app) as client:
        response = client.get('/api/health')
    assert response.status_code == 200
    assert response.json()['status'] == 'ok'


def test_pptx_upload_summarize_and_history_delete(tmp_path):
    with TestClient(app) as client:
        pptx_path = tmp_path / 'graphs.pptx'
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = 'Bipartite Graphs'
        slide.placeholders[1].text = 'A bipartite graph has vertices split into two sets. Matchings pair elements across sets.'
        presentation.save(pptx_path)

        with pptx_path.open('rb') as handle:
            response = client.post('/api/documents', files={'files': ('graphs.pptx', handle, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
        assert response.status_code == 200
        uploaded = response.json()[0]
        assert uploaded['status'] == 'ready'
        assert uploaded['chunks'] >= 1

        response = client.post('/api/summarize', json={'prompt': 'Summarize bipartite graphs'})
        assert response.status_code == 200
        summary = response.json()['summary']
        assert 'Structured Summary' in summary
        assert 'Offline grounded response' not in summary

        history = client.get('/api/history').json()
        assert history
        delete_response = client.delete(f"/api/history/{history[0]['id']}")
        assert delete_response.status_code == 200
